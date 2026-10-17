"""Build the Graduates (UG) role-specific roadmaps from the Word files.

    python scripts/build-ug-role-roadmaps.py <repo> <file1.docx> [<file2.docx> ...]

Every run reads ALL the given Word files and rewrites the outputs, so pass
every batch received so far (1-150.docx, 151-300.docx, ...).

Writes:
  public/roadmaps/ug/<slug>.json            one roadmap per file, loaded by the
                                            report only when a student picks
                                            that career (kept out of page code)
  public/roadmaps/ug/index.json             dropdown role -> roadmap versions, and
                                            what each version was written for
                                            (degree/course), loaded on demand
  data/graduates/ROLE_ROADMAPS_REPORT.md    review sheet: matched, needs review,
                                            duplicates, what was dropped

How a file is read (no heading styles are used, so structure comes from the
text): a paragraph in 16pt+ starts a new roadmap; bold "Step N" / "N." lines
start sections; other short bold lines are sub-headings; bold "Label: text"
lines are notes; Word bullets become lists (in the few files without
bullets, runs of short lines do); "A → B → C" chains and lines joined by "↓"
become a progression; tables keep their header, rows and real links
(tracking parameters removed).

Dropped as not roadmap content: the "Skill-gap engine" section (placeholder
"Not assessed" tables - the report has its own skill-gap page), "Could not
verify" research notes, "Footer:" / "Largest gap:" lines, "↓" separators and
empty lines. "VERIFY" placeholders are reworded into plain guidance.
"""
import json
import os
import re
import sys
from collections import Counter
from difflib import SequenceMatcher, get_close_matches

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

# ---------------------------------------------------------------- text cleanup

LINK_ARTIFACT = re.compile(r"url([^]*)([^\s]*)?")


def clean_url(url):
    url = re.sub(r"[?&]utm_source=chatgpt\.com", "", url or "")
    return url.rstrip("?&")


def clean_text(t):
    t = LINK_ARTIFACT.sub(lambda m: m.group(1), t or "")
    t = re.sub(r"[-]", "", t)          # any other private-use leftovers
    t = t.replace("**", "")
    t = re.sub(r"[ \t ]+", " ", t)
    t = re.sub(r"\s*\n\s*", "\n", t).strip()
    t = re.sub(r"(?i)^for this (specific )?taxonomy entry,\s*(\w)", lambda m: m.group(2).upper(), t)
    # "The key point for this row is ..." / "For this row, X should ..." - the
    # author's spreadsheet row, not something a student knows about.
    t = re.sub(r"(?i)^for this row,\s*(\w)", lambda m: m.group(1).upper(), t)
    t = re.sub(r"(?i)\s+for this row\b", "", t)
    # Notes the author left for the database, not for students.
    if AUTHOR_NOTE.search(t):
        t = " ".join(s for s in re.split(r"(?<=[.;])\s+", t) if not AUTHOR_NOTE.search(s))
    return fix_verify(t)


AUTHOR_NOTE = re.compile(r"(?i)taxonomy-role|role mapping|mapping should|knowledge files?|horizon[- ]file|career-roadmaps|career-clusters|supplied sources|sources? (information )?conflicts?"
                         r"|course[- ](specific |role )?taxonomy|degree taxonomy|internal (degree )?taxonomy|exact taxonomy|taxonomy (entry|match|files?)"
                         r"|roadmap files|(files|taxonomy) could not be (inspected|verified)")


# The author's "VERIFY" marker means different things:
#  - an unknown fee/stipend/date -> real guidance ("check the official site")
#  - a note to the author ("mapping should be VERIFY", "horizon-file output
#    unavailable") -> removed, it is not for students
#  - trailing an ordinary sentence -> just the marker removed
INTERNAL_NOTE = re.compile(r"(?i)taxonomy|knowledge file|horizon[- ]file|horizon source|career-roadmaps|supplied sources|source file|mapping should|career-clusters")
MONEY_CONTEXT = re.compile(r"(?i)fee|cost|price|pay|stipend|paid|amount|deadline|timing|edition|cycle|date|eligib|scholarship|seat|intake|cut-?off")


STRICT_MONEY = re.compile(r"(?i)\bfees?\b|stipend|price|deadline|\bdates?\b|amount|cut-?off|intake|\bseats?\b")


def fix_verify(t):
    if "VERIFY" not in t:
        return t
    if re.search(r"[.;]\s*VERIFY\.?\s*$", t):
        # A bare marker after a finished sentence: guidance only when that
        # sentence is about fees/dates ("cost estimation" is a skill, not a fee).
        t = re.sub(r"\s*VERIFY\.?\s*$", "", t)
        return t + (" Check the official site." if STRICT_MONEY.search(t) else "")
    parts = re.split(r"(?<=[.;])\s+", t)
    t = " ".join(s for s in parts if not ("VERIFY" in s and INTERNAL_NOTE.search(s)))
    t = re.sub(r"^\s*VERIFY\s*[—–-]\s*", "", t)
    if not t.strip():
        return ""
    if t.strip(" .") == "VERIFY":
        return "Check the official site"
    if MONEY_CONTEXT.search(t):
        t = re.sub(r"\bVERIFY current fee\b", "check the current fee on the official site", t)
        t = re.sub(r"\bVERIFY for the relevant cycle\b", "check the official site for the current cycle", t)
        t = re.sub(r"\bVERIFY (for )?(the )?current\b", "check the current", t)
        t = re.sub(r"\bVERIFY because\b", "check the official site, because", t)
        t = re.sub(r"\bVERIFY (the|with|on|against|directly)\b", r"check \1", t)  # "VERIFY the relevant notification"
        t = re.sub(r"\bVERIFY (pay|stipends?|fees?|dates?|timing|eligibility)\b", r"check the current \1 on the official site", t)
        t = re.sub(r"\bVERIFY\b", "check the official site", t)
    else:
        t = re.sub(r"\s*[;:,—–-]?\s*\bVERIFY\b\.?", "", t)
    t = re.sub(r"(\w)\s+check the", r"\1: check the", t)  # "current fee check..." -> "current fee: check..."
    t = re.sub(r"([.!?]\s+)check the", r"\1Check the", t)  # new sentence
    t = re.sub(r"\s{2,}", " ", t).strip()
    return t[:1].upper() + t[1:] if t[:1].islower() and t.startswith("check") else t


def first_link(el, rels):
    """URL of the first hyperlink inside an element (paragraph or cell)."""
    for h in el.iter(W + "hyperlink"):
        rid = h.get(R_ID)
        if rid and rid in rels and "hyperlink" in rels[rid].reltype:
            return clean_url(rels[rid].target_ref)
    m = LINK_ARTIFACT.search("".join(t.text or "" for t in el.iter(W + "t")))
    return clean_url(m.group(2)) if m else None


# ---------------------------------------------------------------- classification

SECTION_KINDS = [
    ("internships", r"internship"),
    ("certifications", r"certification"),
    ("careers", r"hired|careers you|jobs you|career options|job roles"),
    ("pg_india", r"postgraduate|post-graduate|\bpg\b|higher studies"),
    ("abroad", r"abroad"),
    ("growth", r"going forward|progression|career path|long-term|grow"),
    ("build", r"what to build|build the right|skill development|semester|during ug|year-by-year|foundation"),
    ("horizons", r"horizon"),
]
DROP_SECTION = re.compile(r"(?i)skill[- ]gap engine|could not verify")
DROP_LINE = re.compile(r"(?i)^(footer|largest gap)\b[^:]{0,40}:")
# Sub-headings the template uses - recognised even where a file has no bold.
KNOWN_SUBHEADS = {
    "facts", "specialisation", "specialisation options", "specialisation lanes", "government",
    "private and open-source", "government / public-sector routes", "private / open-source routes",
    "first-job titles", "adjacent roles", "application pack", "phd routes", "scholarships and fellowships",
    "scholarships/fellowships", "career horizons", "four career horizons", "skill-gap engine", "could not verify",
    "4-year honours variant", "four-year honours variant", "recommended final skill stack", "target progression",
    "career-route check", "career route check",
}
FACT_KEYS = ("Starts", "Current degree", "Course", "Timeline assumption")
# Facts that are bookkeeping or placeholders, not guidance for the student.
DROP_FACT = re.compile(r"(?i)^(desired career|target career|career cluster|largest gap|footer|verification note|_row_degree|(id|discipline|level|category|typical duration|degree|specialisation|specialization|branch|career id|reference id|reference code|stream|relevant course|education level|reference age|typical entry-level age|broad domain|career category|domain|sector)$)")
# Some files give the facts as a "Career Path" table (ID, Level, Degree,
# Specialization, Target Career, ...) - read into the same fact names.
TABLE_FACT = {"degree": "Current degree", "ug degree": "Current degree", "specialization": "Course", "specialisation": "Course", "branch": "Course",
              "course": "Course", "target career": "Desired career", "career": "Desired career", "career cluster": "Career cluster"}
# Profile tables name the degree and the role many ways: "Recommended Degree",
# "Recommended Qualification", "Undergraduate qualification", "Recommended UG
# Programme"; "Target Job Role", "Target Role".
DEGREE_KEY = re.compile(r"(?i)^((recommended|typical|undergraduate|ug)\s+)*(degree|qualification|programme|program)$")
TARGET_KEY = re.compile(r"(?i)^(target\s+(career|job\s+role|role|job)|desired career)$")


def fact_name(label):
    """A table/line label under the fact name the rest of the build uses."""
    l = label.strip().lower()
    if l in TABLE_FACT:
        return TABLE_FACT[l]
    if DEGREE_KEY.match(l):
        return "Current degree"
    if TARGET_KEY.match(l):
        return "Desired career"
    return label.strip()


def is_profile_table(header, rows):
    """A "Field | Details" table that names the degree or the target role -
    the roadmap's profile ("Career Profile", "Career classification")."""
    hl = [h.strip().lower() for h in header]
    return (len(header) == 2 and hl[0] in ("field", "item", "parameter")
            and any(r and fact_name(r[0]["text"]) in ("Current degree", "Desired career") for r in rows))


def career_path_facts(header, rows):
    """{label: value} when a table is the roadmap's facts, else None."""
    hl = [h.strip().lower() for h in header]
    if any(h == "degree" or h == "ug degree" for h in hl) and len(rows) == 1:
        return {h: c["text"] for h, c in zip(header, rows[0])}
    if is_profile_table(header, rows):
        return {r[0]["text"]: r[1]["text"] for r in rows if len(r) > 1}
    return None


def has_degree_fact(facts):
    return any(fact_name(k) == "Current degree" for k in facts)


def section_kind(heading):
    h = heading.lower()
    for kind, pat in SECTION_KINDS:
        if re.search(pat, h):
            return kind
    return "other"


def strip_heading(h):
    h = re.sub(r"^\d+\.\s*", "", h)  # "2. Step 1 — What to build" -> "What to build"
    h = re.sub(r"^(Step\s*\d+\s*[:—\-–.]?\s*)", "", h, flags=re.I)
    return h.strip() or "Details"


YEAR_BLOCK = re.compile(r"(?i)^(year|semester|sem)\s*\d")
# The standard main headings, written out in full (a sub-heading like "Year 4 /
# Clinical Internship" is not one).
MAIN_HEADING = re.compile(r"(?i)^(internship strategy|internships? (and|&) (field|industry|clinical) (experience|exposure|training)"
                          r"|entry[- ]level career options|careers? you can be hired as|jobs you can target( after .*)?"
                          r"|certifications?|certifications? (and|&) (additional |short )?(learning|courses)|certifications in demand"
                          r"|pg options( after .*)?|postgraduate (study|options)( in india)?|pg (and|&) advanced study( options)?"
                          r"|study abroad( .*)?|career progression|going forward)$")
STANDARD_KINDS = {"build", "internships", "certifications", "careers", "pg_india", "abroad", "growth"}


def role_from_title(title):
    t = clean_text(title)
    t = re.sub(r"(?i)^career\s+roadmap\s*[:·—\-–]\s*", "", t)
    t = re.sub(r"(?i)\s*[:·—\-–]?\s*roadmap$", "", t)
    t = re.sub(r"(?i)^roadmap\s*[:·—\-–]\s*", "", t)
    # "Career Roadmap: Industrial Engineering & Management → Process
    # Improvement Engineer": the course, then the role.
    t = t.split("→")[-1]
    return t.strip(" :·—-–")


def is_bold(p):
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def bold_label(p):
    """('Label', 'rest') when the paragraph opens with a bold 'Label:'."""
    runs = [r for r in p.runs if r.text]
    if not runs or not runs[0].bold:
        return None
    prefix = ""
    for r in runs:
        if not r.bold:
            break
        prefix += r.text
    if ":" not in prefix:
        return None
    k, _, v = clean_text(p.text).partition(":")
    if not k or len(k) > 60:
        return None
    return k.strip(), v.strip()


def font_size(p):
    return next((r.font.size.pt for r in p.runs if r.font.size), None)


# A roadmap starts at its title: "Career roadmap · X", "Career Roadmap: X",
# "X Roadmap", or a large "Career: X". Titles are told apart by their words,
# not only their size - some files set "Step N" headings in the same 24pt as
# titles, and some set titles in plain bold.
TITLE_PAT = re.compile(r"(?i)^(career\s+roadmap\b|roadmap\s*[:·—–-])|\broadmap\s*$")
# The row number some files put in front of a title is not part of it:
# "ID 600 — Career Roadmap: X", and a stray "3Career roadmap · X".
TITLE_PREFIX = re.compile(r"(?i)^(?:id\s*\d+\s*[—–:·-]\s*|\d{1,4}\s*(?=career\s+roadmap\b)|\d{3,4}\s*\|\s*)")  # ... "842 | SAFETY SUPERVISOR — CAREER ROADMAP"
PIPE_ID = re.compile(r"^(\d{3,4})\s*\|")
# "01 | School / Classes 10–12", "14 | COMPLETE ROADMAP": a numbered section.
PIPE_HEADING = re.compile(r"^(\d{1,2})\s*\|\s*(\S.*)$")
# "Career Roadmap — ID 645": the role is only in the facts table that follows.
ID_ONLY = re.compile(r"(?i)^id\s*\d+$")
# "551 — B.Tech / B.E. Software Product Engineering → DevOps Engineer": a
# spreadsheet row header before a roadmap's facts table (and its title).
ROW_HEADER = re.compile(r"^(\d{2,5})\s*[—–-]\s*(.+?)\s*→\s*([^→]+?)\s*$")
# "558. Cloud Support Engineer" / "1656 — Level Designer" followed by its facts
# ("UG: ... Target Career: ..." / "Degree: ... Career: ...") is a roadmap
# title, not a numbered section heading ("20. PG & Advanced Study").
NUM_TITLE = re.compile(r"^(\d{2,5})(?:\.\s+|\s*[—–-]\s*)([^:→]{2,80})$")
FACTS_START = re.compile(r"(?i)^\s*(ug|pg|degree|level|discipline)\s*:")
# "Technical Artist — Career Roadmap": the role first.
ROLE_FIRST_TITLE = re.compile(r"(?i)^(.{2,80}?)\s+[—–-]\s+career\s+roadmap\s*$")
# "UG | Media, Design & Creative Arts | B.Des (Bachelor of Design) | Game Design"
# under such a title: level, domain, degree, course.
PIPE_FACTS = re.compile(r"(?i)^(ug|pg)\s*\|")
# "Career Roadmap 101: ESG Technical Analyst": the number is the pending
# list's Career ID, not part of the title.
NUMBERED_TITLE = re.compile(r"(?i)^(career\s+roadmap)\s+#?(\d{1,4})\s*(?=[:·—–-])")
# The first heading of a roadmap that has no title line: "1. Career Profile",
# "1. Career classification", "1. Career overview". (A "1. ..." list item
# under a progression line is not one.)
PROFILE_HEAD = re.compile(r"(?i)^1\.\s+(career\s+(profile|classification|overview|snapshot|summary|identity)|overview|profile)\b")
# A bare web domain on its own line is an image's source caption, not content.
# Lowercase only: "B.Com" / "M.Com" are degrees, not captions.
SOURCE_CAPTION = re.compile(r"^[a-z0-9-]{3,}(\.[a-z0-9-]+)*\.(com|org|net|in|gov|edu|io|co|uk|br|au|de)(\.[a-z]{2})?/?$")
# "Career ID: 101 · Reference ID: 448 · Discipline: ..." under a title: the
# roadmap's bookkeeping (Career ID = S.No in the pending list, Reference ID =
# the spreadsheet row), never its summary.
CAREER_ID_LINE = re.compile(r"(?i)^career\s*id\s*:\s*\d+")
# Its labels, also when run together with no "·" ("Career ID: 312 Education
# Level: Undergraduate (UG) Career Domain: ... Recommended Degree: ...").
ID_LINE_LABEL = re.compile(r"(?i)(?:^|·|\s)\s*(Career ID|Reference ID|Reference code|Discipline|Education Level|Education Pathway|Career Domain|Career Area|Career category"
                           r"|Recommended Degree|Degree|Specialisation|Specialization|Course|Target Career|Target Role|Level|Stream|Category|Domain|Education"
                           r"|Sector|Qualification|Pathway)\s*:")
# The facts line right under a title, in its other shapes: "Sector: ... |
# Qualification: Diploma in ...", "Career Category: ... Education Pathway: ...
# Target Role: ...", "901 | Manufacturing & Industrial Production".
LEAD_FACTS_LINE = re.compile(r"(?i)^((career\s+id|sector|career\s+category|category|domain|discipline|education(\s+(level|pathway))?|stream|qualification|pathway)\s*:|\d{3,4}\s*\|)")
# Labels whose value is the qualification itself when it reads as one.
ROUTE_LABELS = {"qualification", "pathway", "education pathway", "education"}
# "Career category: ... · Degree: B.Sc. in ..." / "Category: ... · Education:
# ..." right under a title: the roadmap's facts, not its summary.
FACTS_LINE = re.compile(r"(?i)^(career\s+category|category|education(\s+level)?|discipline|domain|stream|level)\s*:.*·.*\b(degree|education|career\s*id)\b")


# An unlabelled degree among the line's pieces ("... · B.Sc – Actuarial Science · ...").
DEGREE_PIECE = re.compile(r"^(B\.\s?[A-Za-z]{1,6}\b|B[A-Z][A-Za-z]{0,5}\b(?=\s|$|\s*[–-])|LL\.?B|MBBS|Diploma\b|ITI\b|D\.\s?Pharm|Bachelor\b)")
# "Career Roadmap • Business, Finance & Entrepreneurship • B.Sc – Actuarial
# Science • Career ID: 862" under a "179 — Pricing Analyst" title.
BULLET_ID_LINE = re.compile(r"(?i)^career\s+roadmap\s*[•·].*\bcareer\s*id\s*:\s*\d+")


def id_line_facts(text):
    """Facts from a Career ID line: labelled values, and an unlabelled degree
    (pieces split by "·", "•" or "|")."""
    parts = ID_LINE_LABEL.split(re.sub(r"\s[|•]\s|•", " · ", text))
    out, loose = {}, parts[0].split("·")
    for label, value in zip(parts[1::2], parts[2::2]):
        bits = [x.strip() for x in value.split("·")]
        if bits[0]:
            name = "Current degree" if label.strip().lower() in ROUTE_LABELS and DEGREE_PIECE.match(bits[0]) else fact_name(label)
            out.setdefault(name, bits[0])
        loose += bits[1:]
    for k, piece in enumerate(x.strip() for x in loose):
        if DEGREE_PIECE.match(piece):
            out.setdefault("Current degree", piece)
        elif k == 0 and re.fullmatch(r"\d{3,4}", piece):
            out.setdefault("ID", piece)  # "901 | Manufacturing & Industrial Production"
    return out


def markdown_table_facts(text):
    """{header: value} from a table typed as text ("| ID | Level | Degree /
    Programme | ... | |---|---|...| | 705 | UG | ..."), else None."""
    if not text.startswith("|") or not re.search(r"\|\s*:?-{3,}", text):
        return None
    cells = [c.strip() for c in text.strip().strip("|").split("|")]
    sep = next((k for k, c in enumerate(cells) if re.fullmatch(r":?-{3,}:?", c)), None)
    if not sep:
        return None
    header = [c for c in cells[:sep] if c]
    values = [c for c in cells[sep:] if c and not re.fullmatch(r":?-{3,}:?", c)]
    return dict(zip(header, values[:len(header)])) if values else None


# Kept in capitals when an ALL-CAPS title or heading is put in normal case.
ACRONYMS = {"ITI", "PG", "UG", "IT", "AI", "CNC", "EV", "HVAC", "COPA", "CAD", "QC", "QA", "GIS", "NCVT", "SCVT", "PMKVY", "NSDC", "DMLT",
            "GNM", "ANM", "IOT", "PLC", "TIG", "MIG", "SCADA", "UAV", "DGCA", "AME", "LPG", "CCTV", "PCB", "MEP", "HR", "OJT", "NCC", "UPSC",
            "SSC", "PV", "MRO", "B.SC", "B.VOC", "D.PHARM", "F&B", "NDT", "PPE", "ICU", "OT", "ECG", "MRI", "CT", "BPO", "KPO", "ERP", "SAP", "SEO", "UI", "UX", "3D", "2D"}


def unshout(s, first_upper=True):
    """ "ABROAD — EDUCATION" -> "Abroad — education", keeping acronyms."""
    words = re.split(r"(\s+)", s.lower())
    out = []
    for w in words:
        core = re.sub(r"[^\w.&+-]", "", w)
        if core.upper() in ACRONYMS or re.fullmatch(r"[a-z]&[a-z]", core):
            w = w.upper()
        elif core in ("india", "indian"):
            w = w.capitalize()
        out.append(w)
    t = "".join(out)
    return (t[:1].upper() + t[1:]) if first_upper else t


def title_case(s):
    """ "SAFETY SUPERVISOR" -> "Safety Supervisor", keeping acronyms."""
    return re.sub(r"[A-Za-z][\w.&'-]*", lambda m: m.group(0).upper() if m.group(0).upper() in ACRONYMS else m.group(0).capitalize(), s.lower())


def next_is_table(els, i):
    """Whether the next non-empty element after els[i] is a table."""
    return next_element(els, i) == "table"


def next_texts(els, i, n):
    """The next n non-empty paragraph texts after els[i] ("table" for a table)."""
    out = []
    for el in els[i + 1:]:
        if el.tag == W + "tbl":
            out.append("table")
        elif el.tag == W + "p":
            t = "".join(x.text or "" for x in el.iter(W + "t")).strip()
            if t:
                out.append(t)
        if len(out) >= n:
            break
    return out


def next_element(els, i):
    """The next non-empty element after els[i]: "table", its text, or ""."""
    for el in els[i + 1:]:
        if el.tag == W + "tbl":
            return "table"
        if el.tag == W + "p":
            t = "".join(x.text or "" for x in el.iter(W + "t")).strip()
            if t:
                return t
    return ""


NOT_TITLE =re.compile(r"(?i)(year|semester|week|month|skills?|learning|study|ug|degree|career|cluster-level|cluster)[\s-]*(by[\s-]*\w+[\s-]*)?roadmap\s*$|roadmap overview|^\d+\.|^\d{1,2}\s*\|")
CAREER_LINE = re.compile(r"(?i)(?:^|·\s*)career:\s*(.+)$")
# The "Degree: ... · Course: ... · Career: ..." line some files put before a
# roadmap's title, and the notes around it that are not for students.
DEGREE_LINE = re.compile(r"(?i)^degree:\s*\S")
PREAMBLE_LINE = re.compile(r"(?i)^(cluster-level roadmap|role-list check\b|\d{1,5}\s*-\s*\d{1,5}$|career roadmap$)")


def title_role(p, text, bold, pending):
    """The role a title line names, or None when the line is not a title."""
    if not text or len(text) > 120 or "\n" in text or re.match(r"(?i)^step\s*\d", text) or DEGREE_LINE.match(text):
        return None
    text = NUMBERED_TITLE.sub(r"\1", TITLE_PREFIX.sub("", text))
    big = (font_size(p) or 0) >= 16
    first = ROLE_FIRST_TITLE.match(text)
    if first and (big or bold) and not re.match(r"(?i)^(year|semester|ug|pg|degree|cluster|skills?|learning|study)\b", first.group(1)):
        return first.group(1).strip()
    looks = bool(TITLE_PAT.search(text)) and not NOT_TITLE.search(text)
    if looks and (big or (bold and (pending or re.match(r"(?i)^career\s+roadmap\s*[:·—–-]", text)))):
        return role_from_title(text)
    m = CAREER_LINE.search(text) if big else None
    return m.group(1).strip(" .") if m else None


META_LABEL = re.compile(r"(?i)(?:^|·|\s)\s*(Degree|Course|Specialisation|Specialization|Target Role|Target Career|Career Area|Career)\s*:")
META_NAME = {"degree": "Current degree", "course": "Course", "specialisation": "Course", "specialization": "Course",
             "career": "Desired career", "target role": "Desired career", "target career": "Desired career"}


def degree_meta(text):
    """{"Current degree", "Course", "Desired career"} from a Degree line -
    "Degree: ... · Course: ... · Career: ...", or the labels run together
    ("Degree: B.Voc – Drone Technology Target Role: ... Career Area: ...")."""
    t = text.replace("\n", " · ")
    parts = META_LABEL.split(t)
    out = {}
    for label, value in zip(parts[1::2], parts[2::2]):
        name, value = META_NAME.get(label.lower()), value.strip(" ·")
        if name and value:
            out.setdefault(name, value)
    return out


def degree_level(s):
    """"doctoral", "pg" or "ug" - a roadmap is only shown to students at the
    level it was written for (a Ph.D. plan starts at Ph.D. Year 1)."""
    s = s or ""
    if re.search(r"(?i)ph\.?\s?d\b|\bd\.\s?sc\b|d\.\s?litt|\bdba\b|doctorate|post-?doctoral|doctoral", s):
        return "doctoral"
    # Integrated / dual-degree programmes (BS-MS, B.Tech+M.Tech, IPM BBA+MBA,
    # LLB+LLM) start straight after Class 12 - undergraduate, though they name
    # a master's. Same rule as INTEGRATED in lib/report/roleRoadmapGrad.tsx.
    if re.search(r"(?i)dual[- ]degree|integrated\s+(programme in management|5-year|five[- ]year|b\.?\s?(tech|sc|ed|a)\b|bs-ms|bba|ll\.?\s?b|m\.?\s?(sc|a|tech)\b|ma\b)|\bipm\b|bba\s*\+\s*mba", s):
        return "ug"
    if re.search(r"(?i)\b(m\.\s?(tech|e|sc|a|com|phil|pharm|des|arch|ed|s)\b|mba\b|pgdm|master|post-?graduate)", s):
        return "pg"
    return "ug"


def name_from_neighbour(rm, prev):
    """For a roadmap pasted without its title: the role numbered right after
    the previous roadmap's ID (as a pending-list S.No, or as a sheet row), if
    this roadmap's own text names it."""
    if not prev or not SHEET_ROWS:
        return None
    m = re.match(r"\s*(\d+)", (prev.get("ids") or {}).get("row", "") or (prev.get("ids") or {}).get("career", ""))
    if not m:
        return None
    n = int(m.group(1)) + 1
    text = " ".join(json.dumps(s["blocks"], ensure_ascii=False) for s in rm["sections"]).lower()
    for row in (PENDING_ROWS.get(n, 0), n):
        if 1 <= row <= len(SHEET_ROWS):
            role = SHEET_ROWS[row - 1][2]
            if re.sub(r"\s*\([^)]*\)", "", role).lower() in text:
                rm["facts"].setdefault("Reference ID", str(row))  # so its degree comes from that row
                return role
    return None


def name_from_careers(rm):
    """The first role a roadmap's "Careers you can be hired as" section names
    ("Ayurvedic Physician, after completing ..." -> "Ayurvedic Physician")."""
    for s in rm["sections"]:
        if s["kind"] != "careers":
            continue
        for blk in s["blocks"]:
            if blk.get("type") == "list" and blk["items"]:
                first = blk["items"][0]["text"]
            elif blk.get("type") == "table" and blk["rows"] and blk["rows"][0]:
                first = blk["rows"][0][0]["text"]
            elif blk.get("type") == "p":
                first = blk["text"]
            else:
                continue  # a sub-heading ("First-job titles"), not a role
            name = re.split(r",| \(| – | - |:", first or "")[0].strip()
            if name and len(name.split()) <= 6:
                return name
    return "Untitled roadmap"


def is_list_item(el):
    """A Word bullet/numbered paragraph - the file's own list marking."""
    return el.find(f"{W}pPr/{W}numPr") is not None


ARROW = ("↓", "→")
ARROW_IN = re.compile(r"[→↓]")


def flow_items(t):
    """The steps of a progression written as a chain - "A → B → C" on one
    line, or lines separated by "↓" / starting with "→" - else None."""
    lines = [x.strip() for x in t.split("\n") if x.strip()]
    if len(lines) > 1:
        if any(x in ARROW for x in lines) or sum(x.startswith("→") for x in lines) >= 2:
            items = [x.lstrip("→↓ ").strip() for x in lines if x not in ARROW]
            return items if len(items) >= 2 else None
        return None
    for sep in ("→", " > "):
        parts = [s.strip() for s in t.split(sep)]
        if len(parts) >= 3 and all(0 < len(s) <= 90 and ". " not in s for s in parts):
            parts[-1] = parts[-1].rstrip(".")
            return parts
    return None


def chain_block(t, label=None):
    """A progression ("A → B → C") or a skill stack ("A + B + C") as a flow
    block, else None."""
    items, sep = flow_items(t), "→"
    if not items:
        parts = [s.strip() for s in t.split(" + ")]
        if "\n" not in t and len(parts) >= 3 and all(0 < len(s) <= 60 and ". " not in s for s in parts):
            items, sep = parts, "+"
    if not items:
        return None
    # "With experience, the pathway can progress toward Senior X → Y": the
    # lead-in becomes the label, the chain starts at "Senior X".
    m = LEAD_IN.match(items[0]) if len(items[0].split()) > 5 else None
    if m and len(m.group(2).split()) <= 6:
        label = f"{label}: {m.group(1)}" if label else m.group(1)
        items = [m.group(2)] + items[1:]
    return {"type": "flow", "items": items, **({"sep": "+"} if sep == "+" else {}), **({"label": label} if label else {})}


LEAD_IN = re.compile(r"^(.*\b(?:towards?|into|to|from|via|is|are|as|be|becomes?|include|includes|:))\s+(\S.*)$")


# ---------------------------------------------------------------- parsing

def parse(path):
    d = docx.Document(path)
    # Which roadmaps mark their lists with Word bullets (almost all). For the
    # few that don't, runs of short lines are read as lists instead. A first
    # pass finds out, so the real pass sees exactly the same roadmaps.
    first = parse_doc(d, path, None)
    roadmaps = parse_doc(d, path, [rm["_li"] for rm in first])
    # A stray section or two pasted with no title (e.g. a lone "04 |
    # Scholarships" between two roadmaps) is not a roadmap of its own.
    keep = []
    for k, rm in enumerate(roadmaps):
        if rm.get("untitled") and len(rm["sections"]) < 3:
            FRAGMENTS.append(f"{os.path.basename(path)}: {' / '.join(s['heading'] for s in rm['sections']) or 'empty'}"
                             f" (after {roadmaps[k - 1]['role'] if k else 'the start'})")
        else:
            keep.append(rm)
    return keep


def parse_doc(d, path, bullets):
    rels = d.part.rels
    roadmaps = []
    cur = None
    section = None
    plain = []  # pending plain lines -> list or paragraphs
    last_num = 0
    uses_steps = False
    sub_num = 0  # > 0 while a numbered list restarted inside a section ("1. Excel", "2. AutoCAD", ...)
    skipping = False  # inside an ignorable part ("Skill-gap engine", "Could not verify")
    pending = None  # facts from a "Degree: ... · Course: ... · Career: ..." line, for the next roadmap
    preamble = False  # between such a line and the title it introduces

    def start_roadmap(title, role):
        nonlocal cur, section, last_num, uses_steps, sub_num, skipping, pending, preamble
        flush_plain()
        if cur is not None and (cur["_n"] == 0 or (not cur["sections"] and not cur["summary"] and norm(role) == norm(cur["role"]))):
            # "Career roadmap · X" then "X Roadmap" (sometimes with its facts
            # table in between): one roadmap, two title lines. A "Degree: ...
            # Career: ..." line between them is this roadmap's own, not the
            # opening of the next one.
            for k, v in (pending or {}).items():
                cur["facts"].setdefault(k, v)
            cur["meta_career"] = cur["meta_career"] or (pending or {}).get("Desired career", "")
            pending, preamble = None, False
            return
        cur = {"title": title, "role": role, "summary": "", "facts": dict(pending or {}), "horizons": [], "sections": [],
               "source": os.path.basename(path), "bullets": bullets[len(roadmaps)] if bullets is not None else True,
               "meta_career": (pending or {}).get("Desired career", ""), "_li": False, "_n": 0}
        roadmaps.append(cur)
        section, last_num, uses_steps, sub_num, skipping = None, 0, False, 0, False
        pending, preamble = None, False

    def item(x):
        return {"text": x["text"], **({"url": x["url"]} if x.get("url") else {})}

    def para_block(x):
        return chain_block(x["text"]) or {"type": "p", **item(x)}

    def flush_plain():
        nonlocal plain
        if not plain or section is None or skipping:
            plain = []
            return
        blocks = section["blocks"]
        run = []

        def end_run(min_len):
            nonlocal run
            if len(run) >= min_len:
                blocks.append({"type": "list", "items": [item(r) for r in run]})
            else:
                blocks.extend(para_block(r) for r in run)
            run = []

        for x in plain:
            if cur["bullets"]:
                in_list = x["li"] and not chain_block(x["text"])
            else:
                # No bullets in this file: short lines run together as a list;
                # long lines, lead-ins ending in ":" and chains stand alone.
                in_list = len(x["text"]) <= 140 and not x["text"].endswith(":") and not chain_block(x["text"])
            if in_list:
                run.append(x)
            else:
                end_run(1 if cur["bullets"] else 2)
                blocks.append(para_block(x))
        end_run(1 if cur["bullets"] else 2)
        plain = []

    def open_section(heading, size):
        """A Step/numbered heading: a new section - except "Year 2 — ..." blocks
        an author numbered at the top level, which stay inside "What to build"."""
        if section is not None and section["kind"] == "build" and YEAR_BLOCK.match(strip_heading(heading)):
            flush_plain()
            section["blocks"].append({"type": "h", "text": strip_heading(heading)})
            return
        # The main headings' size: the largest Step/numbered heading so far.
        cur["_head_size"] = max(cur.get("_head_size", 0), size if size >= 16 else 0)
        new_section(heading)

    def new_section(heading, kind=None):
        nonlocal section, skipping
        flush_plain()
        skipping = bool(DROP_SECTION.search(heading))
        section = {"heading": strip_heading(heading), "kind": kind or section_kind(heading), "blocks": []}
        if DROP_SECTION.search(heading):
            section["drop"] = True
        cur["sections"].append(section)

    def table_rows(el):
        rows = []
        for r in Table(el, d).rows:
            cells = []
            for c in r.cells:
                txt = clean_text(c.text)
                url = first_link(c._tc, rels)
                cells.append({"text": txt, **({"url": url} if url else {})})
            # merged cells repeat; drop exact repeats within a row
            dedup = []
            for c in cells:
                if dedup and c == dedup[-1]:
                    continue
                dedup.append(c)
            rows.append(dedup)
        return rows

    els = list(d.element.body.iterchildren())
    for i, el in enumerate(els):
        if el.tag == W + "p":
            p = Paragraph(el, d)
            text = clean_text(p.text)
            if not text:
                continue
            bold = is_bold(p)
            role = title_role(p, text, bold, pending)
            if role:
                bare_title = TITLE_PREFIX.sub("", text)
                start_roadmap(NUMBERED_TITLE.sub(r"\1", bare_title), role)
                num = NUMBERED_TITLE.match(bare_title) or PIPE_ID.match(text)
                if num:
                    cur["facts"].setdefault("ID", num.group(num.lastindex))  # "Career Roadmap #866: ..." / "842 | SAFETY SUPERVISOR — ..."
                if re.sub(r"[^A-Za-z]", "", role).isupper() and len(re.sub(r"[^A-Za-z]", "", role)) > 3:
                    cur["_caps"] = True  # "SAFETY SUPERVISOR — CAREER ROADMAP": put in normal case later
                left = bare_title.split("→")[0] if "→" in bare_title else ""
                left = re.sub(r"(?i)^career\s+roadmap\s*#?\d*\s*[:·—–-]?\s*", "", left).strip(" :·—–-")
                if left and DEGREE_PIECE.match(left):
                    cur["facts"].setdefault("_row_degree", left)  # "... 670 — B.Sc. Computer Science → IT Support Engineer"
                continue
            num_title = NUM_TITLE.match(text) if bold and not is_list_item(el) else None
            nxt = next_element(els, i) if num_title else ""
            if num_title and FACTS_START.match(nxt) and re.search(r"(?i)career\s*:", nxt):
                start_roadmap(f"Career Roadmap · {num_title.group(2).strip()}", num_title.group(2).strip())
                continue
            if num_title and int(num_title.group(1)) > 25 and (PROFILE_HEAD.match(nxt) or BULLET_ID_LINE.match(nxt) or CAREER_ID_LINE.match(nxt)):
                # "353 — Chip Design Engineer" then straight into "1. Career
                # Profile" (or its Career ID line): a sheet row / S.No and its
                # role, as the title. (A section number is never this large.)
                flush_plain()
                pending = {"ID": num_title.group(1)}
                start_roadmap(f"Career Roadmap · {num_title.group(2).strip()}", num_title.group(2).strip())
                continue
            row = ROW_HEADER.match(text) if (bold or (font_size(p) or 0) >= 16) and not is_list_item(el) else None
            if row and next_is_table(els, i):
                # "551 — B.Tech ... → DevOps Engineer" and its facts table open
                # the next roadmap; its "Career Roadmap · X" title follows, or
                # is missing and the row's career names it.
                flush_plain()
                pending, preamble = {"Desired career": row.group(3)}, "row"
                continue
            if row and re.search(r"[A-Za-z]{3}", row.group(2)):
                # "508 — B.Voc Drone Technology → Agri-Drone Operator" (or "724 —
                # Merchant Navy / Maritime Programmes – ... → Second Officer")
                # then straight into "1. Career classification" - or into a
                # table typed as text, then "1. Career Overview": the row is the
                # title. (A progression "2026–2035 → ..." is never followed by
                # a profile heading.)
                # (A bare "Career Roadmap" line may sit in between, and the
                # profile heading may be unnumbered: "Career profile".)
                nxt_texts = [t for t in next_texts(els, i, 3) if not re.fullmatch(r"(?i)career\s+roadmap", t)]
                first = nxt_texts[0] if nxt_texts else ""
                if (PROFILE_HEAD.match(first) or re.fullmatch(r"(?i)career\s+(profile|overview|classification)", first)
                        or (first.startswith("|") and len(nxt_texts) > 1 and PROFILE_HEAD.match(nxt_texts[1]))):
                    flush_plain()
                    pending = {"Desired career": row.group(3), "_row_degree": row.group(2), "ID": row.group(1)}
                    start_roadmap(f"Career Roadmap · {row.group(3)}", row.group(3))
                    continue
            names_more = re.search(r"(?i)\b(course|career|target role|target career):", text)
            in_facts = section is None or section["kind"] == "facts"
            fresh = cur is not None and cur["_n"] == 0 and not preamble  # just after a title
            if DEGREE_LINE.match(text) and (names_more or (cur is not None and not in_facts)):
                meta = degree_meta(p.text)
                if fresh:
                    # Title first, then "Degree: ... · Course: ... · Career: ...":
                    # this roadmap's own facts.
                    for k, v in meta.items():
                        cur["facts"].setdefault(k, v)
                    cur["meta_career"] = cur["meta_career"] or meta.get("Desired career", "")
                    continue
                # Before the next title (a bare "Degree: B.Tech" inside a Facts
                # block is just a fact).
                flush_plain()
                pending, preamble = meta, True
                continue
            if fresh and PIPE_FACTS.match(text):
                parts = [x.strip() for x in text.split("|")]
                # "UG | domain | degree | course", or "UG | domain | degree":
                # the degree is the part that reads as one, not the domain.
                at = next((k for k, x in enumerate(parts) if k and DEGREE_PIECE.match(x)), None)
                if at is None and len(parts) >= 4:
                    at = len(parts) - 2
                if at is not None:
                    cur["facts"].setdefault("Current degree", parts[at])
                    if at + 1 < len(parts) and parts[at + 1]:
                        cur["facts"].setdefault("Course", parts[at + 1])
                    continue
            if re.match(r"(?i)^(course|career):", text) and (preamble or (fresh and cur["facts"].get("Current degree"))):
                meta = degree_meta(p.text)  # the same line, split over two paragraphs
                if preamble:
                    pending.update(meta)
                else:
                    for k, v in meta.items():
                        cur["facts"].setdefault(k, v)
                    cur["meta_career"] = cur["meta_career"] or meta.get("Desired career", "")
                continue
            if (PREAMBLE_LINE.match(text) and not is_list_item(el)) or re.fullmatch(r"\d{1,3}", text):
                continue  # (a lone "8" / "13" is a stray page or image number; a bulleted "Career roadmap" is a project feature)
            if preamble == "row" and bold_label(p) and not is_list_item(el):
                # "Qualification note: ..." between the row's table and its
                # title: one of the roadmap's facts, not the start of its text.
                k, v = bold_label(p)
                pending.setdefault(k, v)
                continue
            if preamble:
                if preamble != "row" and not (re.match(r"(?i)^(step\s*\d|facts$|career horizons|four career horizons)", text) or text.startswith("Starts:")):
                    continue  # notes between a Degree line and its title
                # The title is missing: the Degree line's (or row's) career names the roadmap.
                start_roadmap(f"{pending.get('Desired career', 'Career')} Roadmap", pending.get("Desired career", "Career"))
            if cur is None:
                continue
            if SOURCE_CAPTION.match(text) and not is_list_item(el):
                continue  # (a bulleted "draw.io" is a tool, not a caption)
            if (CAREER_ID_LINE.match(text) or BULLET_ID_LINE.match(text) or FACTS_LINE.match(text)) and not cur["sections"]:
                for k, v in id_line_facts(text).items():
                    cur["facts"].setdefault(k, v)
                continue
            if not cur["sections"] and not cur["summary"] and not is_list_item(el):
                lead = LEAD_FACTS_LINE.match(text) and (len(id_line_facts(text)) >= 2 or PIPE_ID.match(text))
                # "Engineering, Technology & Computing | ITI Trades – Motor Vehicle ..."
                loose = (len(text) < 260 and re.search(r"\s[|·]\s", text) and not re.search(r"[.!?]\s*$", text)
                         and any(DEGREE_PIECE.match(x.strip()) for x in re.split(r"\s[|·]\s", text)))
                table_text = markdown_table_facts(text)
                if lead or loose:
                    for k, v in id_line_facts(text).items():
                        cur["facts"].setdefault(k, v)
                    continue
                if table_text:
                    for k, v in table_text.items():
                        name = "Current degree" if "degree" in k.lower() or "programme" in k.lower() else fact_name(k)
                        if v.strip():
                            cur["facts"].setdefault(name, v.strip())
                    continue
                sub = re.match(r"^(.{2,80}?)\s+[—–-]\s+(.+)$", text)
                if bold and sub and norm(sub.group(1)) == norm(cur["role"]) and DEGREE_PIECE.match(sub.group(2)):
                    # "IoT Engineer — Bachelor of Science in Electronic Systems" under the title.
                    cur["facts"].setdefault("Current degree", sub.group(2).strip())
                    continue
                if (cur["facts"] and not bold and len(text.split()) <= 4 and len(text) <= 40
                        and not re.search(r"[.!?:]", text) and not ARROW_IN.search(text)):
                    continue  # "Engineering & Technology" / "Electronics": the rest of the facts block
            if (not cur["sections"] and not cur["summary"] and len(text) <= 120 and len(text.split()) <= 14 and DEGREE_PIECE.match(text)
                    and not is_list_item(el) and not re.search(r"[.!?:]\s*$", text)):
                # A bare "B.Sc. – Aviation Science" (or "... → Animator") under the title.
                cur["facts"].setdefault("Current degree", text.split("→")[0].strip(" —–-"))
                continue
            cur["_n"] += 1
            if text in ARROW:
                if plain:
                    plain[-1]["chain"] = True  # the next line continues this progression
                continue
            if DROP_LINE.search(text):
                continue
            li = is_list_item(el)
            if li:
                cur["_li"] = True
            size_p = font_size(p) or 0
            if size_p >= 16:
                bold = True  # a large line that is not a title is a heading
            pipe_head = PIPE_HEADING.match(text) if bold and not li else None
            if pipe_head:
                # "01 | School / Classes 10–12" -> "1. School / Classes 10–12"
                n_pipe = int(pipe_head.group(1))
                text = f"{n_pipe}. {pipe_head.group(2).strip()}"
                if cur.get("_pipe") and last_num >= 10 and n_pipe <= last_num - 5:
                    # "14 | COMPLETE ROADMAP" then "04 | SCHOLARSHIPS ...": the
                    # next roadmap pasted from its 4th section, without a title
                    # (named later from its neighbour's ID and its own text).
                    start_roadmap("", "")
                    cur["_n"], cur["_pipe"] = 1, True
                    last_num = n_pipe - 1
                cur["_pipe"] = True
            step = re.match(r"(?i)^step\s*(\d+)\b", text)
            if not bold and len(text) < 90 and (re.match(r"(?i)^step\s*\d+\s*[:—–-]", text) or (not li and text.lower() in KNOWN_SUBHEADS)):
                bold = True  # files without bold formatting
            num = re.match(r"^(\d+)\.\s+\S", text)
            # Once a roadmap's main headings have a size, a much smaller "Step N"
            # or "N." line is a sub-heading (a project's "Step 1 — Map activities").
            minor = bool(cur.get("_head_size")) and size_p < cur["_head_size"] - 2
            if bold and step and len(text) < 90 and not minor:
                n_step = int(step.group(1))
                seen_steps = cur.setdefault("_steps", set())
                if n_step == 1 and 1 in seen_steps and max(seen_steps) >= 3:
                    # A second Step 1 after a full run of steps, with no title
                    # in between: the next roadmap was pasted without its title
                    # (named later from its own "Careers you can be hired as"
                    # list). A first Step 1 placed late is just out of order.
                    start_roadmap("", "")
                    cur["_n"] = 1
                    seen_steps = cur.setdefault("_steps", set())
                seen_steps.add(n_step)
                uses_steps, sub_num = True, 0
                open_section(text, size_p)
                continue
            if bold and num and not li and len(text) < 90 and not minor:
                n = int(num.group(1))
                # Same size as the roadmap's main headings: a main heading even
                # when the author's numbering skips ("11." then "13.").
                main_size = bool(cur.get("_head_size")) and size_p >= cur["_head_size"] - 2
                if sub_num and n == sub_num + 1 and not main_size:
                    sub_num = n          # still inside a numbered list within this section
                elif n == last_num + 1 or (main_size and n > last_num):
                    last_num, sub_num = n, 0
                    open_section(text, size_p)
                    continue
                else:
                    sub_num = n          # numbering restarted: these are sub-headings
            if (bold and not li and not step and not num and len(text) <= 90 and cur.get("_head_size")
                    and size_p >= cur["_head_size"] - 2):
                # An unnumbered heading at the main headings' size that names a
                # standard section ("Internship Strategy", "PG Options After
                # BCA", "Study Abroad") starts it; others stay sub-headings.
                k = section_kind(text)
                if k in STANDARD_KINDS and (section is None or section["kind"] != k) and not DROP_SECTION.search(text):
                    open_section(text, size_p)
                    continue
            if (bold and not li and not step and not num and not cur.get("_head_size") and section is not None
                    and MAIN_HEADING.match(text)):
                # Files whose headings are all plain bold, no sizes: only the
                # standard main headings ("Internship Strategy", "PG Options",
                # "Career Progression") start a section; others stay sub-headings.
                k = section_kind(text)
                if k in STANDARD_KINDS and section["kind"] != k:
                    open_section(text, size_p)
                    continue

            # A facts paragraph ("UG: ... Specialization: ... Target Career: ...")
            # right under the title is the roadmap's facts, not its summary.
            if section is None and not cur["summary"] and not bold and not (bold_label(p) and re.search(r"(?i)target career:", text)):
                cur["summary"] = text
                continue
            lab = bold_label(p)
            if not lab and not bold and (section is not None or cur["summary"]):
                k, sep, v = text.partition(":")
                if sep and v.strip() and 2 <= len(k) <= 40 and len(k.split()) <= 5 and k[:1].isupper():
                    lab = (k.strip(), v.strip())  # "Starts: 2026", "Eligibility note: ..." in files without bold
            if lab and (section is None or section["kind"] == "facts"):
                # Some files put every fact in one paragraph with line breaks.
                for line in text.split("\n"):
                    k, sep, v = line.partition(":")
                    if sep and k.strip() and len(k) <= 60:
                        cur["facts"].setdefault(k.strip(), v.strip())
                continue
            if bold and (text == "Facts" or text.lower() == "starting point"):  # not "FACTS" (the power-grid technology)
                # "Starting point" is the facts block in some files (Starts,
                # Degree, Target career, ...) - its "Degree:" line is a fact.
                new_section("Facts", "facts")
                continue
            if bold and text.lower() in ("career horizons", "four career horizons", "career horizon"):
                new_section(text, "horizons")
                continue
            if bold and not li and not lab and len(text) <= 90 and section is not None and section["kind"] == "facts":
                # A heading after the facts ("Career-route check") opens real
                # content - it must not be dropped with the facts block.
                new_section(text)
                continue
            if section is None:
                new_section("Overview", "other")
            if lab and len(lab[1]) > 0:
                flush_plain()
                if not skipping:
                    # "0–2 years: Junior X → 2–5 years: Y → ..." is one chain, not a label
                    whole = chain_block(text) if lab[0][:1].isdigit() else None
                    section["blocks"].append(whole or chain_block(lab[1], lab[0]) or {"type": "note", "label": lab[0], "text": lab[1]})
                continue
            if bold and not li and len(text) <= 90:
                flush_plain()
                skipping = bool(DROP_SECTION.search(text))
                if not skipping:
                    # Bold is also used for emphasis: a progression, a skill
                    # stack or a whole sentence is not a sub-heading.
                    section["blocks"].append(chain_block(text) or ({"type": "p", "text": text} if text.endswith((".", "?")) or (text.endswith(":") and len(text) > 40) else {"type": "h", "text": text}))
                continue
            if skipping:
                continue
            if plain and (plain[-1].get("chain") or text.startswith("→")):
                # "A" / "↓" / "B", or "A" / "→ B": one progression
                plain[-1]["text"] += ("\n↓\n" if plain[-1].pop("chain", False) else "\n") + text
                continue
            url = first_link(el, rels)
            plain.append({"text": text, "li": li, **({"url": url} if url else {})})
        elif el.tag == W + "tbl" and preamble == "row":
            # The row header's facts table (ID, Level, Degree, Specialization, Career).
            rows = table_rows(el)
            kv = career_path_facts([c["text"] for c in rows[0]], rows[1:]) if rows else None
            for k, v in (kv or {}).items():
                if v.strip():
                    pending.setdefault(fact_name(k), v.strip())
        elif el.tag == W + "tbl" and cur is not None and not preamble:
            cur["_n"] += 1
            flush_plain()
            rows = table_rows(el)
            if not rows:
                continue
            header = [c["text"] for c in rows[0]]
            hl = " ".join(header).lower()
            if "your level" in hl and "target" in hl:
                continue  # placeholder skill-gap table
            if "period" in hl and "assessment" in hl:
                cur["horizons"] = [{"period": r[0]["text"], "question": r[1]["text"] if len(r) > 1 else "", "output": r[2]["text"] if len(r) > 2 else ""} for r in rows[1:] if r and r[0]["text"]]
                continue
            # The facts table right under the title (a profile table even when a
            # Degree line already gave the degree - setdefault keeps that one).
            early = len(cur["sections"]) <= 1
            kv = career_path_facts(header, rows[1:]) if early and (not has_degree_fact(cur["facts"]) or is_profile_table(header, rows[1:])) else None
            if kv:
                for k, v in kv.items():
                    if v.strip():
                        cur["facts"].setdefault(fact_name(k), v.strip())
                if section is not None and section["blocks"] and section["blocks"][-1]["type"] == "h":
                    section["blocks"].pop()  # its "Career Path" heading
                continue
            if section is None:
                new_section("Overview", "other")
            if skipping:
                continue
            body = [r for r in rows[1:] if any(c["text"] for c in r)]
            section["blocks"].append({"type": "table", "header": header, "rows": body})
    flush_plain()

    for k_rm, rm in enumerate(roadmaps):
        if not rm["role"]:
            # Pasted without its title: the role after its neighbour's in the
            # sheet when its own text names that role, else the first role its
            # careers list names - listed in the review sheet to confirm.
            rm["role"], rm["untitled"] = name_from_neighbour(rm, roadmaps[k_rm - 1] if k_rm else None) or name_from_careers(rm), True
            rm["title"] = f"Career Roadmap · {rm['role']}"
        raw = rm["facts"]
        # Bookkeeping kept for coverage checks only (never written out): the
        # pending list's Career ID and the spreadsheet row ("Reference ID").
        rm["ids"] = {"career": (raw.get("Career ID") or "").strip(), "row": (raw.get("Reference ID") or raw.get("Reference code") or raw.get("ID") or "").strip()}
        # The same facts under other names ("Degree:", "Specialisation:").
        for alias, key in (("Degree", "Current degree"), ("UG", "Current degree"), ("Specialisation", "Course"), ("Specialization", "Course"), ("Branch", "Course"),
                           ("Relevant Course", "Course"), ("Starting point", "Starts"), ("Timeline", "Timeline assumption"), ("Target Career", "Desired career"), ("Target career", "Desired career")):
            if raw.get(alias) and not raw.get(key):
                raw[key] = raw.pop(alias)
        if raw.get("_row_degree") and not raw.get("Current degree"):
            raw["Current degree"] = raw["_row_degree"]  # the row header's "B.Voc Drone Technology"
        if (raw.get("Current degree") and raw.get("Course") and DEGREE_PIECE.match(raw["Course"])
                and not re.search(r"\b(B\.\s?[A-Za-z]{1,6}|B[A-Z][A-Za-z]{0,5}|LL\.?B|MBBS|BA)\b", raw["Current degree"])):
            # "Current degree: Media, Communication, Arts & Design" (the cluster)
            # with "Course: BJMC / B.A. ...": the course line holds the degree.
            raw["Current degree"], raw["Course"] = raw["Course"], ""
        # Labels written for the database, said the way a student reads them.
        for old, new in (("Mapping", "Course alignment"), ("Taxonomy alignment", "Course alignment")):
            if old in raw and new not in raw:
                raw[new] = raw.pop(old)
        rm["desired"] = raw.get("Desired career", "")
        if ID_ONLY.match(rm["role"]) and rm["desired"]:
            # "Career Roadmap — ID 645": named by its facts table's career.
            rm["role"], rm["title"] = rm["desired"], f"Career Roadmap · {rm['desired']}"
        row = sheet_row_for(rm) if SHEET_ROWS else None
        if row:
            # The sheet row it was written for: its exact dropdown name, and its
            # degree and course when the file never states them.
            degree, course, rm["sheet_role"] = SHEET_ROWS[row - 1]
            if not raw.get("Current degree"):
                raw["Current degree"], raw["Course"] = degree, raw.get("Course") or course
        if rm.get("_caps"):
            # "SAFETY SUPERVISOR — CAREER ROADMAP" and "05 | JOB OPTIONS": the
            # sheet's own name for the role, and headings in normal case.
            rm["role"] = rm.get("sheet_role") or title_case(rm["role"])
            rm["title"] = f"Career Roadmap · {rm['role']}"
            for s in rm["sections"]:
                if s["heading"] == s["heading"].upper() and re.search(r"[A-Z]{3}", s["heading"]):
                    s["heading"] = unshout(s["heading"])
                for b in s["blocks"]:
                    if b.get("type") == "h" and b["text"] == b["text"].upper() and re.search(r"[A-Z]{3}", b["text"]):
                        b["text"] = unshout(b["text"])
        # Plain text that sits among the facts (a progression, a note on the
        # route) is real guidance - shown as the roadmap's intro.
        rm["intro"] = [b for s in rm["sections"] if s["kind"] == "facts" and not s.get("drop") for b in s["blocks"]]
        rm["sections"] = [s for s in rm["sections"] if not s.get("drop") and s["kind"] != "facts" and s["blocks"]]
        rm["facts"] = {k: raw[k] for k in FACT_KEYS if raw.get(k)}
        # Other labelled lines in the facts block are real guidance
        # ("Career-route check: ...", "Mapping: ...") - kept as notes.
        rm["about"] = [{"label": k, "text": v} for k, v in raw.items() if v and k not in FACT_KEYS and not DROP_FACT.search(k)
                       and not re.fullmatch(r"(?i)\s*(undergraduate|ug|undergraduate\s*\(ug\))\.?\s*", v)]  # "Education: Undergraduate" says nothing
        # What the roadmap was written for, used to pick between versions of
        # the same role (e.g. Project Engineer - Electrical vs Mechanical).
        # Facts first (the stated degree/course), then the route notes and
        # the opening text, which often name the course ("For an Automobile
        # Engineering graduate, ...").
        first_texts = " ".join(b.get("text", "") for s in rm["sections"][:1] for b in s["blocks"][:3] if b.get("type") == "p")
        if not rm["summary"]:
            firsts = [b["text"] for s in rm["sections"][:2] for b in s["blocks"] if b.get("type") == "p" and len(b["text"]) > 60]
            rm["summary"] = firsts[0] if firsts else ""
        rm["context"] = " ".join([raw.get("Current degree", ""), raw.get("Course", ""), raw.get("Career cluster", ""),
                                  " ".join(a["text"][:160] for a in rm["about"]), rm["summary"][:200], first_texts[:200]]).strip()[:600]
        rm["level"] = degree_level(raw.get("Current degree", "") + " " + raw.get("Course", ""))
    return roadmaps


# ---------------------------------------------------------------- matching

# Reviewed by hand: the file's role name -> the dropdown's name for the same role.
ROLE_ALIASES = {
    "mining surveyor": "Mine Surveyor",
    "meteorologist": "Meteorologist (IMD)",
    "junior design strategist": "Design Strategist (Junior)",
    "junior data analyst": "Data Analyst (Junior)",
    "anganwadi worker": "Aanganwadi Worker",
    "ground duty officer administration logistics": "Ground Duty Officer (Admin/Logistics)",
    "junior engineer mechanical": "Junior Engineer (Mech)",
    # Not aliased on purpose: subject-specific "Mathematics/English/Philosophy
    # Teacher / Lecturer" and "Religious Studies Researcher" - the dropdown's
    # "Teacher/Lecturer" (6 subjects) and "Researcher" (9) are generic, and the
    # version picker would show one subject's plan to students of another.
}


def norm(s):
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def slugify(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:80]


# The sheet's UG rows in order - row N of Roles_Flat_Classified is
# SHEET_ROWS[N - 1] = (degree, course, role). career-clusters.json's UG
# masterRows flatten to exactly these rows (checked against the sheet).
SHEET_ROWS = []
FRAGMENTS = []  # stray untitled pieces skipped by parse_doc, for the report
# Batches numbered by the pending-roadmaps list ("Career ID: 101" = its S.No)
# map S.No -> sheet row through data/graduates/pending-roadmap-rows.json.
PENDING_ROWS = {}


def load_sheet_rows(repo):
    c = json.load(open(os.path.join(repo, "data/graduates/career-clusters.json"), encoding="utf-8"))
    SHEET_ROWS[:] = [(m["degree"], m["course"], r) for m in c["masterRows"] if m.get("level", "UG") == "UG" for r in m["roles"]]
    pending = os.path.join(repo, "data/graduates/pending-roadmap-rows.json")
    if os.path.exists(pending):
        PENDING_ROWS.clear()
        PENDING_ROWS.update({int(k): int(v) for k, v in json.load(open(pending, encoding="utf-8"))["rows"].items()})


def role_tokens(s):
    s = norm(re.sub(r"(?i)\bml\b", "machine learning", re.sub(r"(?i)\bai\b", "artificial intelligence", s or "")))
    return set(s.split()) - {"and", "the", "of", "in"}


def same_role(a, b):
    """The same role written differently: "BMS Engineer" / "Battery Management
    System (BMS) Engineer", "Data Scientist (Junior)" / "Junior Data Scientist"."""
    ta, tb = role_tokens(a), role_tokens(b)
    if bool(ta) and bool(tb) and (ta <= tb or tb <= ta):
        return True
    # one letter apart: "Anganwadi Worker" / "Aanganwadi Worker"
    return bool(a and b) and SequenceMatcher(None, norm(a), norm(b)).ratio() >= 0.94


def sheet_row_for(rm):
    """The sheet row a roadmap was written for: the row its ID names, when that
    row's role is the roadmap's role (an ID is sometimes the pending list's
    S.No and sometimes the sheet row - the role tells them apart)."""
    names = [x for x in (rm["role"], rm.get("desired", ""), rm.get("meta_career", "")) if x]
    for v in (rm["ids"]["row"], rm["ids"]["career"]):
        m = re.match(r"\s*(\d+)", str(v))  # "354 · Engineering, ... · Undergraduate (UG)"
        n = int(m.group(1)) if m else 0
        for row in (n, PENDING_ROWS.get(n, 0)):
            if 1 <= row <= len(SHEET_ROWS) and any(same_role(SHEET_ROWS[row - 1][2], x) for x in names):
                return row
    # No usable ID: the role's only row in the sheet, if it has just one.
    hits = [i + 1 for i, (_, _, r) in enumerate(SHEET_ROWS) if any(norm(r) == norm(x) for x in names)]
    return hits[0] if len(hits) == 1 else None


def name_variants(role):
    """Other ways the dropdown may write a roadmap's role."""
    out = []
    head = re.split(r"\s+[—–]\s+", role)[0]          # "Surveyor — Mining Engineering"
    if head != role:
        out.append(head)
    for r in [role] + out[:]:
        m = re.match(r"(?i)^junior\s+(.+)$", r)        # "Junior Data Scientist"
        if m:
            out.append(f"{m.group(1)} (Junior)")
        m = re.match(r"^(.+?)\s*\(([A-Z][A-Za-z0-9&/]{1,8})\)\s*(.*)$", r)  # "Battery Management System (BMS) Engineer"
        if m and m.group(3):
            out.append(f"{m.group(2)} {m.group(3)}")
    return out


def write_catalog(repo):
    """public/roadmaps/ug/catalog.json: every role with a roadmap, its career
    cluster, and each version's level and what it was written for - what the
    Career Library's "Degree roadmaps" browses (far smaller than index.json)."""
    clusters = json.load(open(os.path.join(repo, "data/graduates/career-clusters.json"), encoding="utf-8"))["clusters"]
    names = [c["name"] for c in clusters]
    cluster_of = {}
    for i, c in enumerate(clusters):
        for r in c["roles"]:
            cluster_of.setdefault(r, i)
    out_dir = os.path.join(repo, "public/roadmaps/ug")
    index = json.load(open(os.path.join(out_dir, "index.json"), encoding="utf-8"))
    short = lambda s: re.sub(r"\s*\([^)]*\)", "", s or "").strip()
    roles = []
    for role in sorted(index["roles"]):
        versions = []
        for slug in index["roles"][role]:
            v = index["versions"].get(slug, {})
            versions.append({"slug": slug, "level": v.get("level", "ug"), "for": short(v.get("for", ""))[:120]})
        roles.append({"role": role, "cluster": cluster_of.get(role, -1), "versions": versions})
    with open(os.path.join(out_dir, "catalog.json"), "w", encoding="utf-8") as fh:
        json.dump({"clusters": names, "roles": roles}, fh, ensure_ascii=False, separators=(",", ":"))
    return len(roles)


def main():
    repo, files = sys.argv[1], sys.argv[2:]
    if files == ["--catalog-only"]:
        # Rebuild just the catalog from the existing index (no Word files needed).
        print(f"catalog: {write_catalog(repo)} roles")
        return
    clusters = json.load(open(os.path.join(repo, "data/graduates/career-clusters.json"), encoding="utf-8"))["clusters"]
    roles = sorted({r for c in clusters for r in c["roles"]})
    load_sheet_rows(repo)
    by_norm = {}
    for r in roles:
        by_norm.setdefault(norm(r), []).append(r)

    all_rm = []
    for f in files:
        all_rm.extend(parse(f))

    out_dir = os.path.join(repo, "public/roadmaps/ug")
    os.makedirs(out_dir, exist_ok=True)
    for name in os.listdir(out_dir):
        if name.endswith(".json"):
            os.remove(os.path.join(out_dir, name))

    index = {"roles": {}, "versions": {}, "roadmaps": []}
    review = []
    slug_count = {}
    seen, duplicates = set(), []
    for rm in all_rm:
        # The same roadmap pasted twice is kept once - same role, opening
        # summary and facts (copies sometimes differ by a stray line further down).
        key = json.dumps([norm(rm["role"]), rm["summary"], rm["facts"], None if rm["summary"] else rm["sections"]], ensure_ascii=False, sort_keys=True)
        if key in seen:
            duplicates.append(rm["title"])
            continue
        seen.add(key)
        base = slugify(rm["role"])
        slug_count[base] = slug_count.get(base, 0) + 1
        slug = base if slug_count[base] == 1 else f"{base}-{slug_count[base]}"
        # The Degree line's "Career:" is the dropdown's own role name, so it
        # wins over a reworded title ("Food R&D Executive" -> "R&D Executive").
        matched, how = by_norm.get(norm(rm.get("meta_career", "")), []), "via the file's 'Degree · Course · Career' line"
        if not matched and rm.get("sheet_role"):
            matched, how = by_norm.get(norm(rm["sheet_role"]), []), "via the sheet row its ID names"
        if not matched:
            matched, how = by_norm.get(norm(rm["role"]), []), "exact"
            if matched and matched != [rm["role"]]:
                how = "same name, different punctuation"
        if not matched:
            # "Associate Product Manager (APM)" -> "Associate Product Manager"
            bare = re.sub(r"\s*\([^)]*\)\s*$", "", rm["role"])
            matched = by_norm.get(norm(bare), []) if bare != rm["role"] else []
            if matched:
                how = "name without the bracketed abbreviation"
        if not matched:
            # "Junior Data Scientist" -> "Data Scientist (Junior)", "Surveyor —
            # Mining Engineering" -> "Surveyor", "Battery Management System
            # (BMS) Engineer" -> "BMS Engineer"
            for v in name_variants(rm["role"]):
                matched = by_norm.get(norm(v), [])
                if matched:
                    how = f"name variant ({v})"
                    break
        if not matched and norm(rm["role"]) in ROLE_ALIASES:
            matched, how = by_norm.get(norm(ROLE_ALIASES[norm(rm["role"])]), []), "reviewed alias (ROLE_ALIASES)"
        if not matched and rm.get("desired"):
            matched = by_norm.get(norm(rm["desired"]), [])
            if matched:
                how = "via the file's 'Desired career' line"
        if rm.get("untitled"):
            how = "title missing in the file - named from its careers list; confirm"
        if matched:
            for r in matched:
                index["roles"].setdefault(r, []).append(slug)
        else:
            how = "needs review"
            review.append((rm["role"], rm["title"], get_close_matches(rm["role"], roles, n=3, cutoff=0.6)))
        # "for": the degree/course the file states it was written for (weighs
        # most when picking a version); "about": the wider text around it.
        # "level": a version is only offered to students at that degree level.
        index["versions"][slug] = {"for": " ".join(v for v in (rm["facts"].get("Current degree", ""), rm["facts"].get("Course", "")) if v),
                                   "course": rm["facts"].get("Course", ""), "about": rm["context"], "level": rm["level"]}
        index["roadmaps"].append({"slug": slug, "title": rm["title"], "role": rm["role"], "mappedTo": matched, "match": how, "source": rm["source"], "level": rm["level"],
                                  "writtenFor": " · ".join(v for v in (rm["facts"].get("Current degree", ""), rm["facts"].get("Course", "")) if v)})
        with open(os.path.join(out_dir, f"{slug}.json"), "w", encoding="utf-8") as fh:
            json.dump({k: rm[k] for k in ("title", "role", "summary", "facts", "about", "intro", "horizons", "sections")}, fh, ensure_ascii=False, separators=(",", ":"))
    multi = {r: s for r, s in index["roles"].items() if len(s) > 1}

    # The report loads this small index on demand (never bundled), then the
    # one roadmap file it needs.
    with open(os.path.join(out_dir, "index.json"), "w", encoding="utf-8") as fh:
        json.dump({"roles": index["roles"], "versions": index["versions"]}, fh, ensure_ascii=False, separators=(",", ":"))
    write_catalog(repo)
    stale = os.path.join(repo, "data/graduates/role-roadmaps-index.json")
    if os.path.exists(stale):
        os.remove(stale)

    lines = ["# UG role roadmaps - mapping review", "",
             f"Built from: {', '.join(os.path.basename(f) for f in files)}", "",
             f"- Roadmaps read: {len(all_rm)}",
             f"- Duplicates skipped (the same roadmap pasted twice; copies may differ by a stray line): {len(duplicates)}" + (f" - {', '.join(duplicates)}" if duplicates else ""),
             f"- Linked to dropdown roles: {sum(1 for r in index['roadmaps'] if r['mappedTo'])} roadmaps -> {len(index['roles'])} dropdown roles",
             f"- Need review (no exact role in the dropdown): {len(review)}",
             f"- Roles with more than one version (the report picks the one closest to the student's course): {len(multi)}",
             "- Written for (degree level): " + ", ".join(f"{lv} {n}" for lv, n in sorted(Counter(r["level"] for r in index["roadmaps"]).items()))
             + " - the UG report only offers a student roadmaps written for their own level (a Ph.D. plan starts at Ph.D. Year 1)", "",
             "## Needs review", "", "| Roadmap role | Title in file | Closest dropdown roles |", "|---|---|---|"]
    for role, title, sugg in review:
        lines.append(f"| {role} | {title} | {', '.join(sugg) or '-'} |")
    untitled = [r for r in index["roadmaps"] if r["match"].startswith("title missing")]
    lines += ["", "## Pasted without a title (named from the roadmap's own careers list - confirm)", ""]
    lines += [f"- {r['role']} ({r['source']})" for r in untitled] or ["- none"]
    lines += ["", "## Stray pieces skipped (a section or two pasted with no title)", ""]
    lines += [f"- {f}" for f in dict.fromkeys(FRAGMENTS)] or ["- none"]
    lines += ["", "## Roles with more than one version", ""]
    by_slug = {r["slug"]: r for r in index["roadmaps"]}
    for role, slugs in multi.items():
        lines.append(f"- **{role}**: " + "; ".join(f"{by_slug[s]['writtenFor'] or 'course not stated'}" for s in slugs))
    if not multi:
        lines.append("- none")
    lines += ["", "## Linked", "", "| Roadmap | Dropdown role(s) | Match |", "|---|---|---|"]
    for r in index["roadmaps"]:
        if r["mappedTo"]:
            lines.append(f"| {r['title']} | {', '.join(r['mappedTo'])} | {r['match']} |")
    with open(os.path.join(repo, "data/graduates/ROLE_ROADMAPS_REPORT.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    print(f"roadmaps {len(all_rm)} | duplicates skipped {len(duplicates)} | linked {sum(1 for r in index['roadmaps'] if r['mappedTo'])} | review {len(review)} | roles with several versions {len(multi)} | dropdown roles covered {len(index['roles'])}")


if __name__ == "__main__":
    main()
